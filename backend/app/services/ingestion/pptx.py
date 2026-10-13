from pptx import Presentation
from typing import List, Dict, Any

def extract_pptx_units(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts slides, bullet points, and speaker notes from PPTX presentation with slide indices.
    """
    prs = Presentation(file_path)
    units = []
    for slide_idx, slide in enumerate(prs.slides):
        slide_num = slide_idx + 1
        slide_title = f"Slide {slide_num}"
        text_runs = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        text_runs.append(line)
            # check if title shape
            if shape == slide.shapes.title and shape.text.strip():
                slide_title = shape.text.strip()

        # Check speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                text_runs.append(f"[Speaker Notes]: {notes}")

        combined_text = "\n".join(text_runs).strip()
        if combined_text:
            units.append({
                "number": slide_num,
                "label": f"Slide {slide_num}",
                "title": slide_title,
                "text": combined_text,
                "topics": [],
                "diagramTitle": f"Slide {slide_num} Visuals",
                "diagramDescription": f"Presentation slide {slide_num} covering {slide_title}.",
                "location": {
                    "slide": slide_num,
                    "section_heading": slide_title
                }
            })
    return units
